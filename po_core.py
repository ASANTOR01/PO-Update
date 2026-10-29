"""
po_core.py  (v2 - supports two buyer PO layouts)
================================================
Core logic for the PO extractor: PDF parsing + Excel export.
Free of any UI code so it can be tested / reused from a script.

Two PDF layouts are supported and auto-detected per file:

  LEGACY  ("WBDCSBC......" POs)
      Customer Order No. | Style No. | Color | XS..XXL | Total | Unit Price | ETD | Ship Mode
      Style heading : "L/S Chase T-Shirt Loose - 36N7.Q16"
      Numbers       : 1,234  /  5.39          (comma thousands, dot decimal)
      ETD           : per row ("15 Oct 26")
      Sub total     : one per page

  BF  ("BF27E....." POs, new format)
      Order No. | Style No. | Color | XS..XXL | Quanity (PCS) | Unit Price (USD)
      Style heading : "Harper Polo, Cotton Jersey, 340 g/m2"  -> Sty + Fabrication
      Numbers       : 1.045  /  16,93         (dot thousands, comma decimal)
      ETD           : once per page header ("Requested ETD 15.02.2027")
      Subtotal      : after every style block; a block may continue on the next page

Why word coordinates instead of plain text?
    Blank size cells (e.g. no XS, or women's styles that stop at L) simply vanish
    from extracted text, so "205 375 305 115" cannot be mapped to sizes by order.
    We therefore read word positions and assign every number to the nearest
    size-header column.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import datetime
from typing import BinaryIO, Optional

import pandas as pd
import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------- #
# Constants
# --------------------------------------------------------------------------- #
SIZES = ["XS", "S", "M", "L", "XL", "XXL"]

# Output columns, matching the latest template ("PO_Summary" dummy sheet).
COLUMNS_FULL = [
    "File", "Sty", "Fabrication", "PO", "Color Code", "Color",
    "XS", "S", "M", "L", "XL", "XXL", "Pc", "Price", "ETD",
]
COLUMNS_NO_FABRICATION = [c for c in COLUMNS_FULL if c != "Fabrication"]

# The template's right-hand block is a set of XLOOKUPs into an external
# workbook. They cannot be derived from a PDF, so only the headers are
# reproduced (cells left empty) so the layout is identical.
LOOKUP_HEADERS = ["Color Code", "Color", "XS", "S", "M", "L", "XL", "XXL", "Pc", "Price", "ETD"]

# Legacy template joins style name and style code with 3 spaces:
#   "L/S Chase T-Shirt Loose   - 36N7.Q16"
STYLE_JOINER = "   "

LAYOUT_LEGACY = "legacy"
LAYOUT_BF = "bf"

# Row-level regexes (layout independent)
RE_ORDER = re.compile(r"^\d{2}[A-Z]{2,}[A-Z]*\d+$")                    # 26TRHK02215 / 26TR007256
RE_STYLE_NO = re.compile(r"^[A-Z0-9]+(?:\.[A-Z0-9]+){2,}$")             # I036893.3IT.GD.03
RE_LEGACY_STYLE_HEADER = re.compile(r"^(?P<name>.+?)\s+-\s+(?P<code>\S+)$")
RE_LEGACY_ETD = re.compile(r"(\d{1,2})\s+([A-Za-z]{3})\s+(\d{2,4})")    # 15 Oct 26
RE_BF_ETD = re.compile(r"Requested\s+ETD\s+(\d{1,2})\.(\d{1,2})\.(\d{4})")


# --------------------------------------------------------------------------- #
# Number formats (the two layouts use opposite separators)
# --------------------------------------------------------------------------- #
@dataclass(frozen=True)
class NumFmt:
    thousands: str
    decimal: str

    @property
    def re_int(self) -> re.Pattern:
        t = re.escape(self.thousands)
        return re.compile(rf"^\d{{1,3}}(?:{t}\d{{3}})+$|^\d+$")

    @property
    def re_dec(self) -> re.Pattern:
        t, d = re.escape(self.thousands), re.escape(self.decimal)
        return re.compile(rf"^\d{{1,3}}(?:{t}\d{{3}})*{d}\d+$|^\d+{d}\d+$")

    def to_int(self, s: str) -> int:
        return int(s.replace(self.thousands, ""))

    def to_float(self, s: str) -> float:
        return float(s.replace(self.thousands, "").replace(self.decimal, "."))


FMT_LEGACY = NumFmt(",", ".")
FMT_BF = NumFmt(".", ",")


# --------------------------------------------------------------------------- #
# Data containers
# --------------------------------------------------------------------------- #
@dataclass
class ParseResult:
    """Result of parsing one PDF."""
    file_label: str
    layout: str = ""
    rows: list[dict] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    printed_total_pcs: Optional[int] = None      # "TOTAL" line on the PO
    printed_total_amount: Optional[float] = None


# --------------------------------------------------------------------------- #
# Generic helpers
# --------------------------------------------------------------------------- #
def file_label_from_name(filename: str) -> str:
    """
    Convert an uploaded filename into the label used in column A.

        WBDCSBC271740_20_.pdf   -> WBDCSBC271740[20]
        WBDCSBC271740 (20).pdf  -> WBDCSBC271740[20]
        WBDCSBC271840.pdf       -> WBDCSBC271840
    """
    stem = re.sub(r"\.pdf$", "", filename.strip(), flags=re.I)
    m = re.match(r"^(.*?)_(\d+)_$", stem) or re.match(r"^(.*?)\s*\((\d+)\)$", stem)
    return f"{m.group(1)}[{m.group(2)}]" if m else stem


def _bf_label(dispatch_no: Optional[str], filename: str) -> str:
    """
    Label for new-format POs: Dispatch No. + the trailing download counter of the
    file name, e.g. 'BF27E02001_LOT_1_UW0287_28__53_.pdf' -> 'BF27E02001[53]'.
    (Edit here if you prefer a different File-column convention.)
    """
    if not dispatch_no:
        return file_label_from_name(filename)
    stem = re.sub(r"\.pdf$", "", filename.strip(), flags=re.I)
    m = re.search(r"[_\s(](\d+)[_)]*$", stem)
    return f"{dispatch_no}[{m.group(1)}]" if m else dispatch_no


def _group_lines(words: list[dict], tol: float = 3.5) -> list[list[dict]]:
    """Group pdfplumber words into visual lines (by vertical position)."""
    lines: list[list[dict]] = []
    for w in sorted(words, key=lambda w: (w["top"], w["x0"])):
        if lines and abs(lines[-1][0]["top"] - w["top"]) <= tol:
            lines[-1].append(w)
        else:
            lines.append([w])
    return [sorted(l, key=lambda w: w["x0"]) for l in lines]


def _center(w: dict) -> float:
    return (w["x0"] + w["x1"]) / 2


def _line_text(line: list[dict]) -> str:
    return " ".join(w["text"] for w in line)


def _is_size_header(line: list[dict]) -> bool:
    texts = [w["text"] for w in line]
    return "XS" in texts and "XXL" in texts


def _size_header_centers(line: list[dict]) -> Optional[tuple[dict, float]]:
    """({size: x_center}, column_pitch) for a header line, else None."""
    centers: dict[str, float] = {}
    for w in line:
        if w["text"] in SIZES and w["text"] not in centers:
            centers[w["text"]] = _center(w)
    if len(centers) != len(SIZES):
        return None
    xs = [centers[s] for s in SIZES]
    return centers, (xs[-1] - xs[0]) / (len(xs) - 1)


def detect_layout(pages_text: list[str]) -> str:
    """Decide which layout a PDF uses from its text."""
    blob = "\n".join(pages_text)
    if "Requested ETD" in blob or "Dispatch No." in blob:
        return LAYOUT_BF
    return LAYOUT_LEGACY


# --------------------------------------------------------------------------- #
# Public entry point
# --------------------------------------------------------------------------- #
def parse_po_pdf(source: BinaryIO | str, filename: str) -> ParseResult:
    """
    Parse one PO PDF (either layout) into item rows.

    `source` may be a path or a binary file-like object (Streamlit upload).
    Only genuine item rows are returned: header blocks, style headings,
    Sub total / TOTAL lines, remarks and signature areas are all ignored.
    """
    with pdfplumber.open(source) as pdf:
        pages = []
        for page in pdf.pages:
            words = page.extract_words(keep_blank_chars=False, use_text_flow=False)
            pages.append(_group_lines(words))

    first = detect_layout([" ".join(_line_text(l) for l in lines) for lines in pages])
    other = LAYOUT_BF if first == LAYOUT_LEGACY else LAYOUT_LEGACY

    result = _parse_pages(pages, first, filename)
    if not result.rows:                       # detection wrong? try the other layout
        alt = _parse_pages(pages, other, filename)
        if alt.rows:
            result = alt
    if not result.rows:
        result.warnings.insert(0, "No item rows found - is this a supported PO layout?")
    return result


# --------------------------------------------------------------------------- #
# Layout-aware page parser
# --------------------------------------------------------------------------- #
def _parse_pages(pages: list[list[list[dict]]], layout: str, filename: str) -> ParseResult:
    is_bf = layout == LAYOUT_BF
    fmt = FMT_BF if is_bf else FMT_LEGACY
    result = ParseResult(file_label=file_label_from_name(filename), layout=layout)

    # State carried across pages: a style block can span a page break, and the
    # size header is re-printed on every page (but we keep the last one as fallback).
    style = fabrication = ""
    last_header = None
    etd: Optional[datetime] = None
    block_pcs, block_amt = 0, 0.0           # running totals since last Subtotal
    dispatch_no = None

    for page_no, lines in enumerate(pages, start=1):
        # ---- page-level fields (BF layout only) ---------------------------
        if is_bf:
            for line in lines:
                text = _line_text(line)
                m = RE_BF_ETD.search(text)
                if m:
                    d, mo, y = map(int, m.groups())
                    try:
                        etd = datetime(y, mo, d)
                    except ValueError:
                        result.warnings.append(f"Page {page_no}: invalid Requested ETD '{m.group(0)}'.")
                if dispatch_no is None and line[0]["text"] == "Dispatch" and len(line) > 2 \
                        and line[1]["text"] == "No.":
                    dispatch_no = line[2]["text"]

        header = next(filter(None, (_size_header_centers(l) for l in lines if _is_size_header(l))), None) \
            or last_header
        if header is None:
            continue                          # page without a table (e.g. terms page)
        last_header = header
        centers, pitch = header
        color_zone_end = centers["XS"] - pitch / 2     # left edge of XS column
        size_zone_end = centers["XXL"] + pitch / 2     # right edge of XXL column

        in_table = False
        for line in lines:
            text = _line_text(line)
            first = line[0]["text"]

            # ---- table boundaries ----------------------------------------
            if _is_size_header(line):
                in_table = True
                continue
            if not in_table:
                continue

            # ---- Sub total / TOTAL (used for reconciliation only) ---------
            if first.lower() in ("sub", "subtotal") or first.upper() == "TOTAL":
                nums = _read_totals(line, fmt)
                if first.upper() == "TOTAL":
                    if nums:
                        result.printed_total_pcs, result.printed_total_amount = nums
                    in_table = False            # nothing after TOTAL is table data
                elif nums:
                    if block_pcs != nums[0] or abs(block_amt - nums[1]) > 1.0:
                        result.warnings.append(
                            f"Page {page_no}: subtotal mismatch - parsed {block_pcs:,} pcs / "
                            f"{block_amt:,.2f} vs printed {nums[0]:,} pcs / {nums[1]:,.2f}.")
                    block_pcs, block_amt = 0, 0.0
                continue
            if first.startswith("Remarks"):
                break

            # ---- item row --------------------------------------------------
            if RE_ORDER.match(first) and len(line) > 1 and RE_STYLE_NO.match(line[1]["text"]):
                row = _parse_item_line(line, fmt, is_bf, centers, color_zone_end,
                                       size_zone_end, etd, result, page_no)
                if row:
                    row["Sty"], row["Fabrication"] = style, fabrication
                    row["File"] = result.file_label
                    result.rows.append(row)
                    block_pcs += row["Pc"]
                    block_amt += row["Pc"] * row["Price"]
                continue

            # ---- style heading --------------------------------------------
            if is_bf:
                # "Harper Polo, Cotton Jersey, 340 g/m2" -> style / fabrication
                name, _, fab = text.partition(",")
                style, fabrication = name.strip(), fab.strip()
            else:
                m = RE_LEGACY_STYLE_HEADER.match(text)
                if m and not any(fmt.re_int.match(w["text"]) for w in line):
                    style = f"{m.group('name')}{STYLE_JOINER}- {m.group('code')}"
                    fabrication = ""

    if is_bf:
        result.file_label = _bf_label(dispatch_no, filename)
        for r in result.rows:
            r["File"] = result.file_label

    _reconcile(result)
    return result


def _read_totals(line: list[dict], fmt: NumFmt) -> Optional[tuple[int, float]]:
    """Pull (pieces, amount) out of a 'Sub total' / 'TOTAL' line (either layout)."""
    toks = [w["text"] for w in line]
    pcs = next((fmt.to_int(t) for t in toks if fmt.re_int.match(t)), None)
    amt = next((fmt.to_float(t) for t in toks if fmt.re_dec.match(t)), None)
    return (pcs, amt) if pcs is not None and amt is not None else None


def _parse_item_line(line, fmt: NumFmt, is_bf: bool, centers, color_zone_end,
                     size_zone_end, etd_page, result: ParseResult, page_no: int) -> Optional[dict]:
    """Turn one visual line of the table into a row dict."""
    order_no, style_no, rest = line[0]["text"], line[1]["text"], line[2:]

    # Colour = everything left of the XS column (digits in a colour name are safe).
    color = " ".join(w["text"] for w in rest if _center(w) < color_zone_end).strip()

    # Sizes: numbers inside the size zone -> nearest header column.
    sizes = {s: 0 for s in SIZES}
    collision = False
    for w in rest:
        if color_zone_end <= _center(w) < size_zone_end and fmt.re_int.match(w["text"]):
            nearest = min(SIZES, key=lambda s: abs(centers[s] - _center(w)))
            collision |= bool(sizes[nearest])
            sizes[nearest] = fmt.to_int(w["text"])
    if collision:
        result.warnings.append(f"{order_no}: two numbers mapped to the same size column (page {page_no}).")

    # Tail columns right of XXL.
    tail = [w["text"] for w in rest if _center(w) >= size_zone_end]
    if is_bf:
        # ["1.000", "PCS", "16,93", "USD"]
        nums = [t for t in tail if t not in ("PCS", "USD")]
        if len(nums) < 2 or not fmt.re_int.match(nums[0]) or not fmt.re_dec.match(nums[1]):
            result.warnings.append(f"{order_no}: could not read Quantity/Unit Price (page {page_no}): {tail}")
            return None
        total, price, etd = fmt.to_int(nums[0]), fmt.to_float(nums[1]), etd_page
        if etd is None:
            result.warnings.append(f"{order_no}: no 'Requested ETD' found on page {page_no}.")
    else:
        # ["377", "5.39", "15", "Oct", "26", "SEA"]
        if len(tail) < 2 or not fmt.re_int.match(tail[0]) or not fmt.re_dec.match(tail[1]):
            result.warnings.append(f"{order_no}: could not read Total/Price columns (page {page_no}): {tail}")
            return None
        total, price = fmt.to_int(tail[0]), fmt.to_float(tail[1])
        etd = _parse_legacy_etd(" ".join(tail[2:5]))
        if etd is None:
            result.warnings.append(f"{order_no}: ETD not recognised in {tail[2:5]}.")

    if sum(sizes.values()) != total:
        result.warnings.append(
            f"{order_no}: size breakdown sums to {sum(sizes.values())} but quantity is {total}.")

    return {"PO": order_no, "Color Code": style_no, "Color": color, **sizes,
            "Pc": total, "Price": price, "ETD": etd}


def _parse_legacy_etd(text: str) -> Optional[datetime]:
    m = RE_LEGACY_ETD.search(text)
    if not m:
        return None
    day, mon, year = m.groups()
    try:
        return datetime.strptime(f"{day} {mon.title()} {year}", "%d %b %y" if len(year) == 2 else "%d %b %Y")
    except ValueError:
        return None


def _reconcile(result: ParseResult) -> None:
    """Compare parsed rows with the TOTAL line printed on the PO."""
    if result.printed_total_pcs is None:
        result.warnings.append("Could not find the TOTAL line to reconcile against.")
        return
    pcs = sum(r["Pc"] for r in result.rows)
    amount = round(sum(r["Pc"] * r["Price"] for r in result.rows), 2)
    if pcs != result.printed_total_pcs:
        result.warnings.append(
            f"Piece count mismatch: parsed {pcs:,} vs printed TOTAL {result.printed_total_pcs:,}.")
    if abs(amount - (result.printed_total_amount or 0)) > 1.0:
        result.warnings.append(
            f"Amount mismatch: parsed {amount:,.2f} vs printed TOTAL {result.printed_total_amount:,.2f}.")


# --------------------------------------------------------------------------- #
# DataFrame + Excel export
# --------------------------------------------------------------------------- #
def build_dataframe(results: list[ParseResult], blank_sizes_as_zero: bool = True,
                    include_fabrication: bool = True, strip_dashes: bool = False) -> pd.DataFrame:
    """
    Combine results into a DataFrame mirroring the template columns.

    blank_sizes_as_zero : False leaves blank PDF size cells empty.
    include_fabrication : False drops the 'Fabrication' column (original 14-col layout).
    strip_dashes        : remove the '---' placeholder the BF layout prints after colours.
    """
    cols = COLUMNS_FULL if include_fabrication else COLUMNS_NO_FABRICATION
    rows = [{c: r.get(c) for c in cols} for res in results for r in res.rows]
    df = pd.DataFrame(rows, columns=cols)
    df = df.replace({"Fabrication": {"": None}})
    if strip_dashes:
        df["Color"] = df["Color"].str.replace(r"\s*---\s*$", "", regex=True).str.strip()
    if not blank_sizes_as_zero:
        for s in SIZES:
            df[s] = df[s].astype("object").where(df[s] != 0, None)
    return df


# Styling copied from the template: (font name, size, bold) per header name.
_HEADER_FONT = {"File": ("Aptos Narrow", 11, False),
                "Sty": ("Arial", 10, True), "Fabrication": ("Arial", 10, True),
                "PO": ("Arial", 10, False),
                "Color Code": ("Arial", 10, True), "Color": ("Arial", 10, True)}
_WIDTHS = {"File": 20.1, "Sty": 38.6, "Fabrication": 28.0, "PO": 12.6, "Color Code": 17.6,
           "Color": 34.6, "XS": 5.9, "S": 5.4, "XL": 6.5, "XXL": 6.9, "Pc": 8.0,
           "Price": 7.4, "ETD": 13.0}
_LOOKUP_WIDTHS = {"Color Code": 17.7, "Color": 35.0, "Pc": 8.0, "ETD": 10.6}


def to_excel_bytes(df: pd.DataFrame, include_lookup_headers: bool = True,
                   pc_as_formula: bool = True) -> bytes:
    """
    Write the DataFrame to an .xlsx reproducing the template: header row, Arial 10
    left-aligned data, bold style column, 'd-mmm-yy' ETD (real dates), '#,##0' piece
    counts, autofilter and column widths.

    pc_as_formula : write Pc as =SUM(XS:XXL) like the template (Excel calculates
                    it on open); False writes the PDF's quantity as a plain number.
    """
    cols = list(df.columns)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    left = Alignment(horizontal="left")

    # Header row: data block, then the (empty) lookup block.
    headers = cols + (LOOKUP_HEADERS if include_lookup_headers else [])
    for i, h in enumerate(headers, start=1):
        in_lookup = i > len(cols)
        name, size, bold = (("Arial", 10, True) if in_lookup and h in ("Color Code", "Color")
                            else ("Arial", 9, False) if in_lookup
                            else _HEADER_FONT.get(h, ("Arial", 9, False)))
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(name=name, size=size, bold=bold)
        c.alignment = left

    xs_col = get_column_letter(cols.index("XS") + 1)
    xxl_col = get_column_letter(cols.index("XXL") + 1)

    for r_idx, rec in enumerate(df.itertuples(index=False), start=2):
        for c_idx, value in enumerate(rec, start=1):
            header = cols[c_idx - 1]
            if isinstance(value, pd.Timestamp):
                value = value.to_pydatetime()
            elif not isinstance(value, str) and pd.isna(value):
                value = None                              # NaN/None -> truly empty cell
            if header == "Pc" and pc_as_formula:
                value = f"=SUM({xs_col}{r_idx}:{xxl_col}{r_idx})"
            c = ws.cell(row=r_idx, column=c_idx, value=value)
            c.font = Font(name="Arial", size=10, bold=(header == "Sty"))
            c.alignment = left
            if header == "ETD":
                c.number_format = "d-mmm-yy"
            elif header == "Pc":
                c.number_format = "#,##0"

    for i, h in enumerate(cols, start=1):
        if h in _WIDTHS:
            ws.column_dimensions[get_column_letter(i)].width = _WIDTHS[h]
    if include_lookup_headers:
        for j, h in enumerate(LOOKUP_HEADERS, start=len(cols) + 1):
            if h in _LOOKUP_WIDTHS:
                ws.column_dimensions[get_column_letter(j)].width = _LOOKUP_WIDTHS[h]
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 2)}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
